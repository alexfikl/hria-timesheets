# SPDX-FileCopyrightText: 2026 Alexandru Fikl <alexfikl@gmail.com>
# SPDX-License-Identifier: MIT

"""Resize the contract columns of the Anexa 1 sheet to fit the number of projects.

The shipped template has 3 contract column pairs (C/D, E/F, G/H) and a total
column (I). This module grows or shrinks that table so the output always shows
exactly the requested number of contracts, then rebuilds every merged range for
the new width.

openpyxl's ``insert_cols`` does not move styles or merges reliably, so we
unmerge everything, move the one column that actually holds content (the total),
clone a column pair for each new contract, and finally re-merge from a known spec.
"""

from __future__ import annotations

import copy

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OLD_PAIRS = 3
OLD_TOTAL_COL = 3 + 2 * OLD_PAIRS  # I (9)
DEFAULT_WIDTH = 8.43

# Footer rows (rows 1-11 are the header block) that are merged across B..last.
FOOTER_MERGED_ROWS = [46, 47, 48, 49, 50, 51, 52, 53, 56, 58, 59]


def _reset(cell: Cell) -> None:
    cell.value = None
    cell.font = Font()
    cell.border = Border()
    cell.fill = PatternFill()
    cell.alignment = Alignment()
    cell.number_format = "General"


def resize(ws: Worksheet, count: int) -> int:
    """Resize ``ws`` for ``count`` contracts and return the total column index."""
    count = max(1, int(count))
    new_total = 3 + 2 * count

    ref_hours_w = ws.column_dimensions["C"].width or 11.57
    ref_interval_w = ws.column_dimensions["D"].width or 12.0
    ref_total_w = ws.column_dimensions[get_column_letter(OLD_TOTAL_COL)].width or 20.57

    max_row = ws.max_row
    for merged in list(ws.merged_cells.ranges):  # ty: ignore[invalid-argument-type]
        ws.unmerge_cells(str(merged))

    # Move the total column (the only column with content to the right) to its
    # new home, then blank the columns it vacated when shrinking.
    if new_total != OLD_TOTAL_COL:
        for row in range(1, max_row + 1):
            src = ws.cell(row, OLD_TOTAL_COL)
            dst = ws.cell(row, new_total)
            dst._style = copy.copy(src._style)
            dst.value = src.value
        if new_total < OLD_TOTAL_COL:
            for col in range(new_total + 1, OLD_TOTAL_COL + 1):
                for row in range(1, max_row + 1):
                    _reset(ws.cell(row, col))
                ws.column_dimensions[get_column_letter(col)].width = DEFAULT_WIDTH

    ws.column_dimensions[get_column_letter(new_total)].width = ref_total_w

    # Clone a contract pair for every contract beyond the template's original 3.
    for index in range(OLD_PAIRS, count):
        hours_col = 3 + 2 * index
        interval_col = 4 + 2 * index
        for row in range(12, 46):
            dst_h = ws.cell(row, hours_col)
            dst_h._style = copy.copy(ws.cell(row, 7)._style)
            dst_h.value = None
            dst_i = ws.cell(row, interval_col)
            dst_i._style = copy.copy(ws.cell(row, 8)._style)
            dst_i.value = None
        ws.cell(13, hours_col).value = "Nr. ore lucrate "
        ws.cell(13, interval_col).value = "Interval orar"
        ws.column_dimensions[get_column_letter(hours_col)].width = ref_hours_w
        ws.column_dimensions[get_column_letter(interval_col)].width = ref_interval_w

    # Rebuild every merge for the new width.
    last = get_column_letter(new_total)
    for row in (2, 3, 4):
        ws.merge_cells(f"B{row}:{last}{row}")
    for row in range(6, 12):
        if new_total >= 6:
            ws.merge_cells(f"B{row}:E{row}")
            ws.merge_cells(f"F{row}:{last}{row}")
        else:
            ws.merge_cells(f"B{row}:C{row}")
            ws.merge_cells(f"D{row}:E{row}")
    ws.merge_cells("B12:B13")
    for index in range(count):
        hours = get_column_letter(3 + 2 * index)
        interval = get_column_letter(4 + 2 * index)
        ws.merge_cells(f"{hours}12:{interval}12")
    ws.merge_cells(f"{last}12:{last}13")
    for row in FOOTER_MERGED_ROWS:
        ws.merge_cells(f"B{row}:{last}{row}")

    return new_total
