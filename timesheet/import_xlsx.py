# SPDX-FileCopyrightText: 2026 Alexandru Fikl <alexfikl@gmail.com>
# SPDX-License-Identifier: MIT

"""Parse a raw PoCIDIF monthly pontaj workbook into Anexa 1 form values.

The imported file is a flat table with one row per contract/project (see
``pontaj_9_2026.xlsx``): identification columns plus ``d1``..``d31`` daily hour
counts. This module turns it into the same string mapping the HTML form submits,
so the regular :func:`web.forms.build_context` parsing and validation still
apply.

The NB (main job) is not encoded in the file: it is always the standard
``8:00-16:00`` day, filled on every working day of the month. Each file row
becomes a project worked after the NB, from ``16:00``, and the rows are stacked
in order when they share a day.
"""

from __future__ import annotations

import calendar
import math
import re
import zipfile
from datetime import date
from io import BytesIO
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from . import intervals
from .holidays import is_weekend, ro_holidays
from .model import EUR_RON_RATE, MAX_CONTRACTS

NB_INTERVAL = "8:00-16:00"
PROJECT_START = 16 * 60  # projects start right after the NB, at 16:00
MAX_DAY = 31
MAX_ROWS = 1000  # guard against workbooks whose max_row is huge but mostly empty
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


class PontajImportError(ValueError):
    """Raised when an uploaded workbook cannot be understood."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        result = float(str(value).replace(",", "."))
    except ValueError:
        return None
    return result if math.isfinite(result) else None


def _format_number(value: float) -> str:
    return f"{value:g}"


def _header_row(ws: Worksheet) -> tuple[int, dict[str, int]]:
    """Find the header row and return its ``{lowercased name: column}`` map."""
    limit = min(ws.max_row or 1, 10)
    for row in range(1, limit + 1):
        names: dict[str, int] = {}
        for col in range(1, (ws.max_column or 1) + 1):
            name = re.sub(r"[\s_]+", "", _text(ws.cell(row, col).value).lower())
            if name:
                names.setdefault(name, col)
        if "domeniu" in names and "nume" in names:
            return row, names
    raise PontajImportError(
        "Fișierul nu conține un tabel de pontaj recunoscut "  # spell: disable
        "(lipsesc coloanele «Nume»/«Domeniu»)."
    )


def _select_sheet(wb: Workbook) -> Worksheet:
    if "Pontaj" in wb.sheetnames:
        return wb["Pontaj"]
    return wb.active


def _read_records(
    ws: Worksheet, header: int, names: dict[str, int]
) -> list[dict[str, Any]]:
    day_columns = {
        day: names[f"d{day}"] for day in range(1, MAX_DAY + 1) if f"d{day}" in names
    }

    def cell(row: int, key: str) -> Any:
        col = names.get(key)
        return ws.cell(row, col).value if col else None

    records: list[dict[str, Any]] = []
    last_row = min(ws.max_row or header, header + MAX_ROWS)
    for row in range(header + 1, last_row + 1):
        domeniu = _text(cell(row, "domeniu"))
        nume = _text(cell(row, "nume"))
        if not domeniu and not nume:
            continue
        records.append(
            {
                "domeniu": domeniu,
                "nume": nume,
                "functie": _text(cell(row, "functie")),
                "ore_cim": _number(cell(row, "orecim")),
                "tarif": _number(cell(row, "tarif")),
                "an": _number(cell(row, "an")),
                "luna": _number(cell(row, "luna")),
                "hours": {
                    day: _number(ws.cell(row, col).value)
                    for day, col in day_columns.items()
                },
            }
        )
    if not records:
        raise PontajImportError("Fișierul nu conține niciun rând de pontaj.")
    return records


def parse_pontaj(data: bytes) -> dict[str, str]:
    """Return form values parsed from a pontaj workbook.

    Raises :class:`PontajImportError` (a :class:`ValueError`) when the file is
    empty, is not a valid workbook, or has no usable rows.
    """
    if not data:
        raise PontajImportError("Fișierul încărcat este gol.")

    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            uncompressed = sum(info.file_size for info in zf.infolist())
            if uncompressed > MAX_UNCOMPRESSED_BYTES:
                raise PontajImportError("Fișierul este prea mare la decomprimare.")  # spell: disable
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True)
    except PontajImportError:
        raise
    except Exception as exc:  # openpyxl and zipfile raise several unrelated types
        raise PontajImportError("Fișierul nu este un document Excel valid.") from exc

    ws = _select_sheet(wb)
    header, names = _header_row(ws)
    records = _read_records(ws, header, names)
    if len(records) >= MAX_CONTRACTS:
        raise PontajImportError(
            f"Fișierul conține {len(records)} proiecte "  # spell: disable
            f"(maximul permis este {MAX_CONTRACTS - 1})."  # spell: disable
        )

    # Identification comes from the POCIDIF contract row (falling back to the
    # first row if none is present).
    main = next(
        (r for r in records if r["domeniu"].upper().startswith("POCIDIF")),
        records[0],
    )

    today = date.today()
    year = int(main["an"]) if main["an"] else today.year
    month = int(main["luna"]) if main["luna"] else today.month
    # Match the range build_context() accepts so calendar.monthrange() is safe.
    year = max(2000, min(2100, year))
    month = max(1, min(12, month))

    parts = main["nume"].title().split()
    family_name = parts[0] if parts else ""
    given_name = " ".join(parts[1:])
    function = main["functie"].split("/", 1)[0].strip().title()

    values: dict[str, str] = {
        "family_name": family_name,
        "given_name": given_name,
        "function": function,
        "year": str(year),
        "month": str(month),
        "count": str(1 + len(records)),
    }
    if main["tarif"]:
        values["ron_rate"] = _format_number(main["tarif"])
        values["euro_rate"] = f"{main['tarif'] / EUR_RON_RATE:.2f}"

    # The NB is the standard 8:00-16:00 day on every working day of the month.
    days_in_month = calendar.monthrange(year, month)[1]
    holidays = ro_holidays(year)
    for day in range(1, days_in_month + 1):
        current = date(year, month, day)
        if not is_weekend(current) and current not in holidays:
            values[f"interval_{day}_0"] = NB_INTERVAL

    # Each file row is a project worked from 16:00, stacked per day in order.
    next_start: dict[int, int] = {}
    for offset, record in enumerate(records):
        index = offset + 1  # contract 0 is the NB
        values[f"project_{index}"] = record["domeniu"]
        if record["ore_cim"]:
            values[f"max_hours_{index}"] = _format_number(record["ore_cim"])

        for day, hours in record["hours"].items():
            if hours is None or hours <= 0:
                continue
            start = next_start.get(day, PROJECT_START)
            end = min(start + round(hours * 60), intervals.MINUTES_PER_DAY)
            if end <= start:
                continue
            values[f"interval_{day}_{index}"] = intervals.format(
                [intervals.TimeRange(start, end)]
            )
            next_start[day] = end

    return values
