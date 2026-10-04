# SPDX-FileCopyrightText: 2026 Alexandru Fikl <alexfikl@gmail.com>
# SPDX-License-Identifier: MIT

"""Fill a copy of the official Anexa 1 template with a :class:`Timesheet`.

The template is shipped blanked (no personal data) with 31 day rows; see
``tools/build_template.py`` for how it was produced. Contract columns are
resized per request in :mod:`timesheet.layout`.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from .holidays import is_weekend, ro_holidays
from .layout import resize

TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent / "templates" / "anexa1_blank.xlsx"
)

MONTHS_RO = {
    1: "IANUARIE",
    2: "FEBRUARIE",
    3: "MARTIE",
    4: "APRILIE",
    5: "MAI",
    6: "IUNIE",
    7: "IULIE",
    8: "AUGUST",
    9: "SEPTEMBRIE",
    10: "OCTOMBRIE",
    11: "NOIEMBRIE",
    12: "DECEMBRIE",
}

DAY_FIRST_ROW = 14
LAST_DAY_ROW = 44  # day 31
TOTAL_ROW = 45

WEEKEND_FILL = PatternFill(fill_type="solid", start_color="F2F2F2", end_color="F2F2F2")
HOLIDAY_FILL = PatternFill(fill_type="solid", start_color="FCE4D6", end_color="FCE4D6")
INACTIVE_FILL = PatternFill(fill_type="solid", start_color="D9D9D9", end_color="D9D9D9")


def _maybe_number(value: str):
    text = str(value).strip()
    return int(text) if text.isdigit() else text


def _set_text(cell, value) -> None:
    """Write text as a literal string so Excel never evaluates it as a formula."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _value_column(count: int) -> int:
    """Column holding the personal fields in rows 6-11 (F, or D for one contract)."""
    return 6 if count >= 2 else 4


def render(ts, gray: bool = True, write_date: bool = True) -> bytes:
    """Return the generated workbook as bytes."""
    count = len(ts.projects)
    if count < 1:
        raise ValueError("At least one contract is required.")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb["Pontaj"]

    total_col = resize(ws, count)
    value_col = _value_column(count)

    ws.cell(4, 2).value = f"Luna {MONTHS_RO[ts.month]} Anul {ts.year}"

    personal = [
        ts.display_name,
        _maybe_number(ts.cnp) if ts.cnp else None,
        ts.function or None,
        ts.euro_rate or None,
        ts.ron_rate or None,
        ts.leader or None,
    ]
    for offset, value in enumerate(personal):
        _set_text(ws.cell(6 + offset, value_col), value)

    for index, project in enumerate(ts.projects):
        _set_text(ws.cell(12, 3 + 2 * index), project.header)

    days = ts.days_in_month
    for day in range(1, days + 1):
        row = DAY_FIRST_ROW + day - 1
        ws.cell(row, 2).value = day
        for index in range(count):
            cell = ts.cells.get((day, index))
            hours_col = 3 + 2 * index
            interval_col = 4 + 2 * index
            ws.cell(row, hours_col).value = (
                float(cell.hours) if cell and cell.hours else None
            )
            _set_text(
                ws.cell(row, interval_col),
                (cell.interval or "").strip() or None if cell else None,
            )
        ws.cell(
            row, total_col
        ).value = f"=SUM(C{row}:{get_column_letter(total_col - 1)}{row})"

    # Blank out day rows beyond the current month (e.g. 31 in a 30-day month).
    for day in range(days + 1, 32):
        row = DAY_FIRST_ROW + day - 1
        ws.cell(row, 2).value = None
        for col in range(3, total_col + 1):
            ws.cell(row, col).value = None

    for index in range(count):
        letter = get_column_letter(3 + 2 * index)
        ws.cell(
            TOTAL_ROW, 3 + 2 * index
        ).value = f"=SUM({letter}{DAY_FIRST_ROW}:{letter}{LAST_DAY_ROW})"
    total_letter = get_column_letter(total_col)
    ws.cell(
        TOTAL_ROW, total_col
    ).value = f"=SUM({total_letter}{DAY_FIRST_ROW}:{total_letter}{LAST_DAY_ROW})"

    _set_text(
        ws["B47"],
        f"Numele şi prenumele persoană: {ts.display_name}" if ts.display_name else None,
    )
    if write_date:
        last_day = date(ts.year, ts.month, days)
        _set_text(
            ws["B49"],
            f"Data: {last_day.day:02d}.{last_day.month:02d}.{last_day.year}",
        )
    _set_text(
        ws["B51"],
        f"Luat la cunoștință de către Responsabil Lider/Partener: {ts.leader}"
        if ts.leader
        else None,
    )

    if gray:
        holidays = ro_holidays(ts.year)
        for day in range(1, 32):
            row = DAY_FIRST_ROW + day - 1
            if day > days:
                fill = INACTIVE_FILL
            else:
                current = date(ts.year, ts.month, day)
                if current in holidays:
                    fill = HOLIDAY_FILL
                elif is_weekend(current):
                    fill = WEEKEND_FILL
                else:
                    fill = None
            if fill is None:
                continue
            for col in range(2, total_col + 1):
                ws.cell(row, col).fill = fill

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
